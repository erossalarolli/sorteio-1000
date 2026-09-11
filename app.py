import os
from datetime import datetime
from zoneinfo import ZoneInfo

from flask import Flask, render_template, request, redirect, url_for, session, jsonify, flash, send_file
import psycopg
from psycopg.rows import dict_row
from psycopg.errors import UniqueViolation
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "troque-esta-chave")
DATABASE_URL = os.environ.get("DATABASE_URL")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123")

def db():
    if not DATABASE_URL:
        raise RuntimeError("DATABASE_URL não configurada.")
    return psycopg.connect(DATABASE_URL, row_factory=dict_row)

def init_db():
    # Cada número reservado é uma linha própria (id próprio), mesmo quando
    # vários números são escolhidos juntos na mesma reserva. Isso facilita
    # editar/excluir um número específico sem afetar os outros.
    with db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS reservas (
                id BIGSERIAL PRIMARY KEY,
                numero INTEGER UNIQUE NOT NULL CHECK (numero BETWEEN 1 AND 1000),
                nome TEXT NOT NULL,
                telefone TEXT NOT NULL,
                cpf TEXT,
                criado_em TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)
        conn.commit()

    # Migração: se o projeto passou pela versão que agrupava números em
    # "pedidos", traz de volta qualquer número que só exista lá para a
    # tabela "reservas" (uma linha por número), sem duplicar nem perder nada.
    with db() as conn:
        existe_pedidos = conn.execute(
            "SELECT to_regclass('public.numeros_reservados') AS t"
        ).fetchone()
        if existe_pedidos and existe_pedidos["t"]:
            faltantes = conn.execute("""
                SELECT nr.numero, p.nome, p.telefone, p.cpf, p.criado_em
                FROM numeros_reservados nr
                JOIN pedidos p ON p.id = nr.pedido_id
                WHERE NOT EXISTS (
                    SELECT 1 FROM reservas r WHERE r.numero = nr.numero
                )
            """).fetchall()
            for r in faltantes:
                conn.execute("""
                    INSERT INTO reservas (numero, nome, telefone, cpf, criado_em)
                    VALUES (%s, %s, %s, %s, %s)
                    ON CONFLICT (numero) DO NOTHING
                """, (r["numero"], r["nome"], r["telefone"], r["cpf"], r["criado_em"]))
            if faltantes:
                conn.commit()

@app.before_request
def ensure_db():
    try:
        init_db()
    except Exception:
        app.logger.exception("Falha ao migrar/inicializar o banco de dados.")

def normalizar_telefone(valor):
    return "".join(ch for ch in (valor or "") if ch.isdigit())

def normalizar_cpf(valor):
    return "".join(ch for ch in (valor or "") if ch.isdigit())

def cpf_valido(cpf):
    if not cpf or len(cpf) != 11:
        return False
    if cpf == cpf[0] * 11:
        return False

    def calc_digito(base):
        soma = sum(int(d) * peso for d, peso in zip(base, range(len(base) + 1, 1, -1)))
        resto = (soma * 10) % 11
        return 0 if resto == 10 else resto

    d1 = calc_digito(cpf[:9])
    d2 = calc_digito(cpf[:9] + str(d1))
    return cpf[-2:] == f"{d1}{d2}"

def parse_numeros(valor):
    """Aceita string 'CSV' (ex: '1,2,3') ou lista, e devolve uma lista de inteiros únicos."""
    if isinstance(valor, list):
        itens = valor
    else:
        itens = str(valor or "").split(",")
    numeros = set()
    for item in itens:
        texto = str(item).strip()
        if texto == "":
            continue
        numeros.add(int(texto))
    return sorted(numeros)

@app.get("/")
def index():
    with db() as conn:
        rows = conn.execute(
            "SELECT numero FROM reservas ORDER BY numero"
        ).fetchall()
    ocupados = [r["numero"] for r in rows]
    return render_template("index.html", ocupados=ocupados)

@app.post("/reservar")
def reservar():
    nome = (request.form.get("nome") or "").strip()
    telefone = normalizar_telefone(request.form.get("telefone"))
    cpf = normalizar_cpf(request.form.get("cpf"))

    if not nome or len(nome) < 2:
        flash("Informe seu nome.")
        return redirect(url_for("index"))

    if len(telefone) < 10 or len(telefone) > 13:
        flash("Informe um WhatsApp/telefone válido.")
        return redirect(url_for("index"))

    if not cpf_valido(cpf):
        flash("Informe um CPF válido.")
        return redirect(url_for("index"))

    try:
        numeros = parse_numeros(request.form.get("numeros"))
    except ValueError:
        flash("Número inválido.")
        return redirect(url_for("index"))

    if not numeros:
        flash("Escolha pelo menos um número.")
        return redirect(url_for("index"))

    if any(not (1 <= n <= 1000) for n in numeros):
        flash("Um dos números escolhidos é inválido.")
        return redirect(url_for("index"))

    try:
        with db() as conn:
            placeholders = ", ".join(["(%s, %s, %s, %s)"] * len(numeros))
            params = []
            for n in numeros:
                params.extend([n, nome, telefone, cpf])

            inseridos = conn.execute(f"""
                INSERT INTO reservas (numero, nome, telefone, cpf)
                VALUES {placeholders}
                ON CONFLICT (numero) DO NOTHING
                RETURNING numero
            """, params).fetchall()

            if len(inseridos) != len(numeros):
                conn.rollback()
                inseridos_set = {r["numero"] for r in inseridos}
                conflitantes = [n for n in numeros if n not in inseridos_set]
                lista = ", ".join(f"{n:04d}" for n in conflitantes)
                flash(f"O(s) número(s) {lista} já foi(ram) reservado(s) por outra pessoa. Escolha outro(s).")
            else:
                conn.commit()
                lista = ", ".join(f"{n:04d}" for n in numeros)
                flash(f"Número(s) {lista} reservado(s) com sucesso!")

    except Exception:
        flash("Não foi possível concluir a reserva. Tente novamente.")

    return redirect(url_for("index"))

@app.get("/admin/login")
def admin_login():
    if session.get("admin"):
        return redirect(url_for("admin"))
    return render_template("admin_login.html")

@app.post("/admin/login")
def admin_login_post():
    senha = request.form.get("senha", "")
    if senha == ADMIN_PASSWORD:
        session["admin"] = True
        return redirect(url_for("admin"))
    flash("Senha incorreta.")
    return redirect(url_for("admin_login"))

@app.get("/admin/logout")
def admin_logout():
    session.pop("admin", None)
    return redirect(url_for("index"))

def admin_required():
    return session.get("admin") is True

@app.get("/admin")
def admin():
    if not admin_required():
        return redirect(url_for("admin_login"))

    with db() as conn:
        rows = conn.execute("""
            SELECT id, numero, nome, telefone, cpf, criado_em,
                   CASE WHEN cpf IS NOT NULL AND cpf <> ''
                        THEN COUNT(*) OVER (PARTITION BY cpf)
                        ELSE 1
                   END AS qtd_cpf
            FROM reservas
            ORDER BY criado_em ASC, id ASC
        """).fetchall()

    tz = ZoneInfo("America/Cuiaba")
    reservas = []
    for r in rows:
        dt = r["criado_em"]
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=ZoneInfo("UTC"))
        r["data_hora"] = dt.astimezone(tz).strftime("%d/%m/%Y %H:%M:%S")
        reservas.append(r)

    return render_template("admin.html", reservas=reservas)

@app.get("/api/reservas")
def api_reservas():
    if not admin_required():
        return jsonify({"erro": "não autorizado"}), 401
    with db() as conn:
        rows = conn.execute("""
            SELECT id, numero, nome, telefone, cpf, criado_em
            FROM reservas
            ORDER BY criado_em ASC, id ASC
        """).fetchall()
    return jsonify(rows)

@app.post("/admin/reservas/<int:reserva_id>/excluir")
def admin_excluir_reserva(reserva_id):
    if not admin_required():
        return jsonify({"erro": "não autorizado"}), 401
    with db() as conn:
        row = conn.execute(
            "DELETE FROM reservas WHERE id = %s RETURNING id",
            (reserva_id,)
        ).fetchone()
        conn.commit()
    if not row:
        return jsonify({"erro": "reserva não encontrada"}), 404
    return jsonify({"ok": True})

@app.post("/admin/reservas/<int:reserva_id>/editar")
def admin_editar_reserva(reserva_id):
    if not admin_required():
        return jsonify({"erro": "não autorizado"}), 401

    data = request.get_json(silent=True) or request.form
    nome = (data.get("nome") or "").strip()
    telefone = normalizar_telefone(data.get("telefone"))
    cpf = normalizar_cpf(data.get("cpf"))
    numero_raw = data.get("numero")

    try:
        numero = int(numero_raw)
    except (TypeError, ValueError):
        return jsonify({"erro": "Número inválido."}), 400

    if not nome or len(nome) < 2:
        return jsonify({"erro": "Informe o nome."}), 400
    if len(telefone) < 10 or len(telefone) > 13:
        return jsonify({"erro": "Informe um WhatsApp/telefone válido."}), 400
    if not cpf_valido(cpf):
        return jsonify({"erro": "Informe um CPF válido."}), 400
    if not 1 <= numero <= 1000:
        return jsonify({"erro": "Número deve ser entre 1 e 1000."}), 400

    try:
        with db() as conn:
            row = conn.execute("""
                UPDATE reservas
                SET numero = %s, nome = %s, telefone = %s, cpf = %s
                WHERE id = %s
                RETURNING id
            """, (numero, nome, telefone, cpf, reserva_id)).fetchone()
            conn.commit()
    except UniqueViolation:
        return jsonify({"erro": "Esse número já está em uso por outra reserva."}), 409
    except Exception:
        return jsonify({"erro": "Não foi possível salvar as alterações."}), 500

    if not row:
        return jsonify({"erro": "reserva não encontrada"}), 404
    return jsonify({"ok": True})

@app.post("/admin/reservas/zerar")
def admin_zerar_reservas():
    if not admin_required():
        return jsonify({"erro": "não autorizado"}), 401
    with db() as conn:
        conn.execute("DELETE FROM reservas")
        conn.commit()
    return jsonify({"ok": True})

@app.get("/admin/exportar")
def admin_exportar_excel():
    if not admin_required():
        return redirect(url_for("admin_login"))

    with db() as conn:
        rows = conn.execute("""
            SELECT numero, nome, telefone, cpf, criado_em
            FROM reservas
            ORDER BY numero ASC
        """).fetchall()

    tz = ZoneInfo("America/Cuiaba")

    wb = Workbook()
    ws = wb.active
    ws.title = "Reservas"

    cabecalho = ["Número", "Nome", "WhatsApp", "CPF", "Data/hora"]
    ws.append(cabecalho)
    for col in range(1, len(cabecalho) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0B2530", end_color="0B2530", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        dt = r["criado_em"]
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=ZoneInfo("UTC"))
        data_hora = dt.astimezone(tz).strftime("%d/%m/%Y %H:%M:%S")
        cpf_fmt = r["cpf"] or ""
        if cpf_fmt and len(cpf_fmt) == 11:
            cpf_fmt = f"{cpf_fmt[:3]}.{cpf_fmt[3:6]}.{cpf_fmt[6:9]}-{cpf_fmt[9:]}"
        ws.append([f"{r['numero']:04d}", r["nome"], r["telefone"], cpf_fmt, data_hora])

    larguras = [12, 30, 18, 18, 20]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[chr(64 + i)].width = largura

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"sorteio_1000_reservas_{datetime.now(tz).strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nome_arquivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.get("/health")
def health():
    return "OK", 200

@app.post("/sortear")
def sortear():
    if not admin_required():
        return jsonify({"erro": "não autorizado"}), 401

    with db() as conn:
        row = conn.execute("""
            SELECT numero, nome, telefone
            FROM reservas
            ORDER BY random()
            LIMIT 1
        """).fetchone()

    if not row:
        return jsonify({"erro": "Ainda não há números reservados para sortear."}), 400

    return jsonify({
        "numero": f"{row['numero']:04d}",
        "nome": row["nome"],
        "telefone": row["telefone"]
    })

if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
